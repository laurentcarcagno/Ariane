"""Excel export of an analysed harness (built on the server, from the stored model)."""
from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date
from typing import Literal
from functools import partial
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

# Cover sheet shared with REQMAP (HMA-C-049 template). Only its HEADER sheet is kept.
HEADER_TEMPLATE = Path(__file__).resolve().parents[3] / "config" / "templates" / "HMA-C-049_header.xlsx"

_CLASSIFICATION_CELLS = {"public": "B10", "commercial": "C10", "internal": "D10"}
_Text = Field(default="", max_length=2000)


class ExportSettings(BaseModel):
    """What the PDFs cannot tell: entered by the user before exporting. Client, reference and date are deduced."""

    edition: str = Field(default="00", max_length=20)
    subtitle: str = _Text
    classification: Literal["", "public", "commercial", "internal"] = ""
    export_control: Literal["", "Not Listed/Not Controlled", "Listed/Controlled"] = ""
    quality_manager: str = _Text
    project_manager: str = _Text
    redactor: str = _Text


COVER_SHEET = "Page de garde"
HEAD_FILL = PatternFill("solid", fgColor="12293A")
HEAD_FONT = Font(bold=True, color="FFFFFF")
SEVERITY_FILL = {"error": PatternFill("solid", fgColor="FBE9E9"), "warning": PatternFill("solid", fgColor="FBF3DC")}


BANNER_ROWS = 4
HEAD_ROW = BANNER_ROWS + 2  # one blank row between the banner and the table
GUTTER = 2  # empty columns left of the tables: they hold the logo, so the banner looks the same whatever the table widths
GUTTER_WIDTH = 14  # characters: 2 x 103 px
LOGO_PX = (200, 61)  # the REQMAP logo (509x156) scaled to fit the gutter


@dataclass
class Banner:
    """The 4-line header of the data sheets: logo, title, reference, date, edition."""

    logo: bytes | None
    ref: str
    date: str
    edition: str


def _banner(ws, banner: Banner, title: str) -> None:
    for r in range(1, BANNER_ROWS + 1):
        ws.row_dimensions[r].height = 20
    for col in range(1, GUTTER + 1):
        ws.column_dimensions[get_column_letter(col)].width = GUTTER_WIDTH
    if banner.logo:
        img = XLImage(BytesIO(banner.logo))
        img.width, img.height = LOGO_PX
        ws.add_image(img, "A1")
    for r, (text, bold) in enumerate([(title, True), (f"Réf : {banner.ref}", False), (f"Date : {banner.date}", False),
                                      (f"Édition : {banner.edition}", False)], 1):
        c = ws.cell(r, GUTTER + 1, text)  # starts right of the gutter and overflows over the empty cells
        if text.startswith("="):
            c.data_type = "s"
        c.font = Font(bold=bold, size=14 if bold else 10, color="12293A")
        c.alignment = Alignment(vertical="center")


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list], fills: list[PatternFill | None] | None = None, *,
           banner: Banner) -> None:
    ws = wb.create_sheet(title)
    pad = [None] * GUTTER
    for _ in range(HEAD_ROW - 1):
        ws.append([None])
    ws.append(pad + head)
    for c in ws[HEAD_ROW][GUTTER:]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for i, r in enumerate(rows):
        ws.append(pad + r)
        if fills and fills[i]:
            for c in ws[ws.max_row][GUTTER:]:
                c.fill = fills[i]
    for i, h in enumerate(head, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None]) + 2
        ws.column_dimensions[get_column_letter(i + GUTTER)].width = min(width, 70)
    _banner(ws, banner, f"Harness {banner.ref} — {title}")
    ws.freeze_panes = f"A{HEAD_ROW + 1}"
    ws.auto_filter.ref = f"{get_column_letter(GUTTER + 1)}{HEAD_ROW}:{get_column_letter(GUTTER + len(head))}{ws.max_row}"
    for row in ws.iter_rows(min_row=HEAD_ROW + 1):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=False)


def _text(ws, ref: str, value: str) -> None:
    ws[ref] = value
    if value.startswith("="):  # never let user or extracted text become a formula
        ws[ref].data_type = "s"


def _cover(model: dict, customer: str, username: str, st: ExportSettings) -> tuple[Workbook, bytes | None]:
    """The template's HEADER sheet, filled from the model and the settings; empty workbook without the template."""
    if not HEADER_TEMPLATE.exists():
        wb = Workbook()
        wb.remove(wb.active)
        return wb, None
    wb = load_workbook(HEADER_TEMPLATE)
    logo = wb["REQUIREMENTS"]._images[0]._data()  # the wide logo the template uses on its data sheets
    for name in wb.sheetnames:
        if name != "HEADER":
            del wb[name]
    for name in list(wb.defined_names):  # they pointed into the template's sheets, which are gone or renamed
        del wb.defined_names[name]
    ws = wb["HEADER"]
    ws.title = COVER_SHEET
    ws.sheet_view.view = "normal"  # the template opens in page-break preview, which prints a large "Page 1" over the sheet
    pn = model.get("partNumber") or ""
    _text(ws, "H2", customer)
    _text(ws, "H3", pn)
    _text(ws, "H9", pn)
    ws["H10"] = date.today().strftime("%d/%m/%Y")
    _text(ws, "H11", st.edition.strip() or "00")
    _text(ws, "B16", f"Harness From-To analysis {pn}".strip())
    _text(ws, "B19", st.subtitle.strip() or f"Generated by ARIANE ({model.get('plugin', '')})")
    if st.classification:
        for key, ref in _CLASSIFICATION_CELLS.items():
            cell = ws[ref]
            cell.value = f"{'☒' if key == st.classification else '☐'} {str(cell.value).lstrip('☒☐ ')}"
            if key == st.classification:
                font = copy(cell.font)
                font.b = True
                cell.font = font
    for ref, name, role in (("D37", st.quality_manager, "Quality Manager"), ("D42", st.project_manager, "Project Manager"),
                            ("D47", st.redactor.strip() or username, "")):
        if name.strip():
            _text(ws, ref, "\n".join(p for p in (name.strip(), role) if p))
    if st.export_control:
        ws["B53"] = f"EXPORT CONTROL STATUS: {st.export_control}"
    return wb, logo


def build_workbook(model: dict, customer: str = "", username: str = "", settings: ExportSettings | None = None) -> bytes:
    st = settings or ExportSettings()
    wb, logo = _cover(model, customer, username, st)
    banner = Banner(logo, model.get("partNumber") or "", date.today().strftime("%d/%m/%Y"), st.edition.strip() or "00")
    sheet = partial(_sheet, wb, banner=banner)
    flagged = {(r.get("conn"), r.get("pin")) for i in model["issues"] if i["severity"] == "error" for r in i["refs"] if r.get("pin")}
    segs = {s["id"]: s for s in (model.get("topology") or {}).get("segments", [])}

    rows, fills = [], []
    for w in sorted(model["wires"], key=lambda w: w["conductorId"]):
        to = w.get("to") or {}
        route = model["routes"].get(w["pairId"], [])
        diam = [segs[s]["diameterMm"] for s in route if s in segs]
        rows.append([w["conductorId"], w["pairId"], w["polarity"], w["signal"], w["nomRed"], w.get("type"),
                     w["from"]["conn"], w["from"]["pin"], to.get("conn"), to.get("pin"),
                     " > ".join(route) if route else "NON ROUTÉ", max((d for d in diam if d), default=None)])
        bad = (w["from"]["conn"], w["from"]["pin"]) in flagged or (to.get("conn"), to.get("pin")) in flagged or not route
        fills.append(SEVERITY_FILL["error"] if bad else None)
    sheet("From-To", ["Fil", "Paire", "Polarité", "Signal", "Nom/Red", "Type", "De connecteur", "De pin", "Vers connecteur",
                           "Vers pin", "Segments", "Ø max (mm)"], rows, fills)

    rows, fills = [], []
    for c in model["connectors"]:
        for k in c["contacts"]:
            bad = (c["id"], k["pin"]) in flagged
            rows.append([c["ref"], c["partNumber"], k["pin"], ", ".join(k["wireIds"]), k.get("origin") or k.get("dest"), k.get("wire")])
            fills.append(SEVERITY_FILL["error"] if bad else None)
    sheet("Connecteurs", ["Connecteur", "Référence", "Pin", "Fils", "Origine / destination (plan)", "N° fil (plan)"], rows, fills)

    sheet("Segments", ["Segment", "Nœud A", "Nœud B", "Ø (mm)", "Repères nomenclature", "Longueur (pt plan)"],
           [[s["id"], s["a"], s["b"], s["diameterMm"], " / ".join(s["bomItems"]), s["lengthPt"]] for s in segs.values()])
    sheet("Nomenclature", ["Pos", "Référence", "Description", "Quantité", "Unité", "Notes", "TBD"],
           [[b["position"], b["partNumber"], b["description"], b["amount"], b["unit"], b["notes"], "oui" if b["tbd"] else ""] for b in model["bom"]])
    order = {"error": 0, "warning": 1, "info": 2}
    issues = sorted(model["issues"], key=lambda i: order[i["severity"]])
    sheet("Rapport", ["Gravité", "Code", "Message"], [[i["severity"], i["code"], i["message"]] for i in issues],
           [SEVERITY_FILL.get(i["severity"]) for i in issues])

    buf = BytesIO()
    wb.active = 0
    wb.save(buf)
    return buf.getvalue()
