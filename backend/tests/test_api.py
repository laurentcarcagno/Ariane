from pathlib import Path

import pytest

from tests.conftest import make_client

ROOT = Path(__file__).resolve().parents[2]
DRW, PL = ROOT / "DRW3500402-907_Init_Lower_1-6.pdf", ROOT / "PL_3500402-907_A.1.pdf"
have_samples = pytest.mark.skipif(not (DRW.exists() and PL.exists()), reason="sample PDFs missing")


def bearer(keys, **kw):
    return {"Authorization": f"Bearer {keys.token(**kw)}"}


def upload(client, headers=None):
    with open(DRW, "rb") as d, open(PL, "rb") as b:
        return client.post("/api/v1/imports", data={"plugin": "beyond_gravity"}, files={"drawing": d, "bom": b}, headers=headers or {})


def test_should_expose_public_config_without_secret(client):
    r = client.get("/api/v1/config")
    assert r.status_code == 200 and set(r.json()) == {"authDisabled", "authority", "clientId"}


def test_should_require_a_token(client):
    assert client.get("/api/v1/plugins").status_code == 401
    assert client.get("/api/v1/plugins", headers={"Authorization": "Bearer garbage"}).status_code == 401


def test_should_reject_expired_wrong_audience_and_missing_customer(client, keys):
    for kw in ({"exp": 1}, {"aud": "other"}, {"iss": "https://evil/realms/x"}):
        assert client.get("/api/v1/me", headers=bearer(keys, **kw)).status_code == 401
    assert client.get("/api/v1/me", headers=bearer(keys, customer=None)).status_code == 403


def test_should_list_only_plugins_enabled_for_the_customer(client, keys):
    r = client.get("/api/v1/plugins", headers=bearer(keys))
    assert [p["id"] for p in r.json()] == ["beyond_gravity"]
    assert client.get("/api/v1/plugins", headers=bearer(keys, customer="unknown-corp")).status_code == 403


def test_should_reject_non_pdf_upload(dev_client):
    r = dev_client.post("/api/v1/imports", data={"plugin": "beyond_gravity"}, files={"drawing": ("a.pdf", b"nope"), "bom": ("b.pdf", b"nope")})
    assert r.status_code == 422


@have_samples
def test_should_import_store_and_reopen(client, keys):
    h = bearer(keys)
    r = upload(client, h)
    assert r.status_code == 200
    body = r.json()
    assert len(body["model"]["wires"]) == 96 and len(body["model"]["routes"]) == 48
    listing = client.get("/api/v1/imports", headers=h).json()
    assert [i["id"] for i in listing] == [body["id"]] and listing[0]["errors"] == 4
    assert client.get(f"/api/v1/imports/{body['id']}", headers=h).json()["model"]["partNumber"] == "3500402-907"


@have_samples
def test_should_not_show_another_customers_import(tmp_path, keys):
    """Isolation: a second configured customer sees nothing of the first, and gets 404 (not 403) by id."""
    import json
    cfg = tmp_path / "customers.json"
    cfg.write_text(json.dumps({"customers": {"beyond-gravity": {"label": "BG", "plugins": ["beyond_gravity"]},
                                              "other": {"label": "Other", "plugins": ["beyond_gravity"]}}}))
    with make_client(tmp_path, keys, customers_config_path=cfg) as c:
        mine = upload(c, bearer(keys, customer="beyond-gravity")).json()["id"]
        theirs = bearer(keys, customer="other")
        assert c.get("/api/v1/imports", headers=theirs).json() == []
        assert c.get(f"/api/v1/imports/{mine}", headers=theirs).status_code == 404
        assert c.delete(f"/api/v1/imports/{mine}", headers=theirs).status_code == 404
        assert c.get(f"/api/v1/imports/{mine}", headers=bearer(keys, customer="beyond-gravity")).status_code == 200
        assert c.delete(f"/api/v1/imports/{mine}", headers=bearer(keys, customer="beyond-gravity")).status_code == 204


def test_should_refuse_plugin_not_enabled_for_customer(tmp_path, keys):
    import json
    cfg = tmp_path / "customers.json"
    cfg.write_text(json.dumps({"customers": {"acme": {"label": "Acme", "plugins": []}}}))
    with make_client(tmp_path, keys, customers_config_path=cfg) as c:
        r = c.post("/api/v1/imports", data={"plugin": "beyond_gravity"}, files={"drawing": ("a.pdf", b"%PDF"), "bom": ("b.pdf", b"%PDF")},
                   headers=bearer(keys, customer="acme"))
        assert r.status_code == 403


@have_samples
def test_should_export_xlsx_for_the_owner_only(tmp_path, keys):
    import io
    import json

    from openpyxl import load_workbook
    cfg = tmp_path / "customers.json"
    cfg.write_text(json.dumps({"customers": {"beyond-gravity": {"label": "BG", "plugins": ["beyond_gravity"]},
                                              "other": {"label": "O", "plugins": ["beyond_gravity"]}}}))
    with make_client(tmp_path, keys, customers_config_path=cfg) as c:
        mine = bearer(keys, customer="beyond-gravity")
        imp = upload(c, mine).json()["id"]
        r = c.get(f"/api/v1/imports/{imp}/export.xlsx", headers=mine)
        assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Page de garde", "From-To", "Connecteurs", "Segments", "Nomenclature", "Rapport"]
        assert wb["Page de garde"]["H2"].value == "BG" and wb["Page de garde"]["H9"].value == "3500402-907"
        ft = wb["From-To"]
        assert ft.max_row == 102 and ft["C7"].value == "907-01+" and ft["C6"].value == "Fil" and ft["A5"].value is None
        banner = [c.value for r in ft.iter_rows(min_row=1, max_row=4) for c in r if c.value]
        assert banner[1:] == ["Réf : 3500402-907", f"Date : {__import__('datetime').date.today():%d/%m/%Y}", "Édition : 00"] and len(ft._images) == 1
        assert all(len(wb[n]._images) == 1 for n in wb.sheetnames[1:])
        assert {wb[n]["C2"].value for n in wb.sheetnames[1:]} == {"Réf : 3500402-907"}  # same banner position on every sheet
        assert c.get(f"/api/v1/imports/{imp}/export.xlsx", headers=bearer(keys, customer="other")).status_code == 404
        r = c.post(f"/api/v1/imports/{imp}/export.xlsx", headers=mine,
                   json={"edition": "02", "classification": "internal", "quality_manager": "=1+1", "export_control": "Listed/Controlled"})
        head = load_workbook(io.BytesIO(r.content))["Page de garde"]
        assert head["H11"].value == "02" and head["D10"].value.startswith("☒") and head["C10"].value.startswith("☐")
        assert head["D37"].value.startswith("=1+1") and head["D37"].data_type == "s" and "Listed/Controlled" in head["B53"].value
        assert head.sheet_view.view == "normal"
        assert c.post(f"/api/v1/imports/{imp}/export.xlsx", headers=mine, json={"classification": "x"}).status_code == 422
        assert c.post(f"/api/v1/imports/{imp}/export.xlsx", headers=bearer(keys, customer="other"), json={}).status_code == 404


def test_should_serve_requests_without_any_oidc_setting_when_auth_is_disabled(tmp_path, keys):
    """Regression: the dev setup (no issuer, no JWKS URL) must not try to build a JWKS client."""
    s = make_client(tmp_path, keys, auth_disabled=True, oidc_issuer="")
    s.app.state.key_resolver = None
    with s as c:
        assert c.get("/api/v1/me").json()["customer"] == "beyond-gravity"
        assert c.get("/api/v1/plugins").status_code == 200
